"""Local business logic. No remote calls or unverified accounting integrations."""
from __future__ import annotations
import calendar
import csv
import hashlib
import io
import json
import math
import random
import sqlite3
import statistics as st
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

def now():
    return datetime.now().astimezone().isoformat(timespec='seconds')

def year_start(d):
    return d.replace(year=d.year-1, day=min(d.day, calendar.monthrange(d.year-1,d.month)[1]))

SCHEMAS = {
    'points': ['id','name'],
    'items': ['id','name','kind','unit','price','cost','pack','target','fallback'],
    'sales': ['date','point','dish','qty','available','closed'],
    'batches': ['id','point','ingredient','qty','unit','usable','checked_at','use_before'],
    'recipes': ['dish','ingredient','amount','unit','version','valid_from'],
    'deliveries': ['id','point','ingredient','qty','unit','arrival','confirmed'],
    'movements': ['id','point','ingredient','date','kind','qty','unit','reason'],
    'schedules': ['point','next_delivery','interval_days'],
}
LABELS = {'points':'Торговые точки','items':'Ассортимент','sales':'Продажи закрытых смен',
          'batches':'Остатки по партиям','recipes':'Технологические карты',
          'deliveries':'Поставки в пути','movements':'Движения (журнал)','schedules':'График поставок'}

class DataError(ValueError): pass

class AccountingAdapter:
    """Contract for a future, separately agreed integration."""
    def read(self, path):
        raise NotImplementedError

class FileAdapter(AccountingAdapter):
    def __init__(self): self.is_demo=False
    def read(self, path, kind=None):
        path=Path(path)
        if path.suffix.lower()=='.xlsx':
            from openpyxl import load_workbook
            wb=load_workbook(path, read_only=True, data_only=True)
            result={}
            try:
                if '_meta' in wb.sheetnames:
                    self.is_demo=wb['_meta']['B1'].value=='DEMO'
                for ws in wb:
                    if ws.title not in SCHEMAS: continue
                    rows=ws.iter_rows(values_only=True)
                    headers=[str(x or '').strip() for x in next(rows, [])]
                    check_headers(ws.title,headers)
                    result[ws.title]=[{k: normalize(v) for k,v in zip(headers,row)}
                                     for row in rows if any(x is not None for x in row)]
            finally: wb.close()
            if not result: raise DataError('В XLSX нет известных листов. Используйте шаблон.')
            return result
        if path.suffix.lower()!='.csv': raise DataError('Поддерживаются XLSX и CSV.')
        kind=kind or path.stem
        if kind not in SCHEMAS: raise DataError('Выберите тип таблицы CSV.')
        raw=path.read_bytes()
        try: txt=raw.decode('utf-8-sig')
        except UnicodeDecodeError: txt=raw.decode('cp1251')
        delim=';' if txt.partition('\n')[0].count(';')>txt.partition('\n')[0].count(',') else ','
        reader=csv.DictReader(io.StringIO(txt), delimiter=delim)
        check_headers(kind,reader.fieldnames or [])
        records=list(reader)
        self.is_demo=bool(records) and all(r.get('_dataset')=='DEMO' for r in records)
        return {kind:records}

def check_headers(kind,headers):
    missing=set(SCHEMAS[kind])-set(headers)
    if missing: raise DataError(f'{kind}: нет столбцов '+', '.join(sorted(missing)))
    if len(headers)!=len(set(headers)): raise DataError(f'{kind}: повторяются имена столбцов')

def normalize(v):
    if isinstance(v,datetime): return v.date().isoformat()
    if isinstance(v,date): return v.isoformat()
    return '' if v is None else v

def number(v, label, nonnegative=True):
    if v is None or str(v).strip()=='': raise DataError(f'{label}: значение не заполнено')
    try: x=float(str(v).replace(',','.'))
    except (ValueError,TypeError): raise DataError(f'{label}: требуется число')
    if not math.isfinite(x) or (nonnegative and x<0): raise DataError(f'{label}: некорректное число')
    return x

def isodate(v, label):
    try: return date.fromisoformat(str(v)).isoformat()
    except (ValueError,TypeError): raise DataError(f'{label}: нужна дата ГГГГ-ММ-ДД')

def flag(v, label):
    if str(v).lower() not in ('0','1','true','false'): raise DataError(f'{label}: только 0 или 1')
    return str(v).lower() in ('1','true')

def validate(tables, existing=None):
    existing=existing or {}
    merged={k:tables.get(k,existing.get(k,[])) for k in SCHEMAS}
    errors=[]
    for kind,rows in tables.items():
        if kind not in SCHEMAS: errors.append(f'Неизвестная таблица {kind}'); continue
        seen=set()
        for i,r in enumerate(rows,2):
            try:
                for field in SCHEMAS[kind]:
                    if field not in r: raise DataError(f'нет столбца {field}')
                    if field!='use_before' and (r[field] is None or str(r[field]).strip()==''): raise DataError(f'не заполнено {field}')
                key=tuple(str(r[x]).strip() for x in ({'sales':['date','point','dish'],'recipes':['dish','ingredient','version','valid_from'],'schedules':['point']}.get(kind,['id'])))
                if key in seen: raise DataError(f'дубликат ключа {key}')
                seen.add(key)
                for f in ('qty','amount','price','cost','pack','target','fallback','interval_days'):
                    if f in r:
                        v=number(r[f],f)
                        if f in ('amount','pack','interval_days') and v<=0: raise DataError(f'{f} должен быть > 0')
                        if f=='interval_days' and (v!=int(v) or v>60): raise DataError('interval_days: целое от 1 до 60')
                        r[f]=v
                for f in ('date','valid_from','arrival','next_delivery'):
                    if f in r: isodate(r[f],f)
                if r.get('use_before'): isodate(r['use_before'],'use_before')
                if 'checked_at' in r:
                    try: datetime.fromisoformat(str(r['checked_at']))
                    except ValueError: raise DataError('checked_at: нужна дата и время ISO')
                for f in ('available','closed','usable','confirmed'):
                    if f in r: r[f]=int(flag(r[f],f))
                if kind=='sales' and not flag(r['closed'],'closed'): raise DataError('незакрытая смена: не является подтвержденной продажей')
                if kind=='sales' and date.fromisoformat(str(r['date']))>=date.today(): raise DataError('допускаются только закрытые смены до сегодняшнего дня')
                if kind=='items':
                    if r['kind'] not in ('dish','ingredient'): raise DataError('kind: dish или ingredient')
                    if r['unit'] not in ('шт','л','кг'): raise DataError('единица: шт, л или кг')
                    if r['kind']=='dish' and r['unit']!='шт': raise DataError('блюда учитываются в шт')
                if kind=='movements' and r['kind'] not in ('receipt','consumption','writeoff','adjustment'): raise DataError('неизвестный вид движения')
                for f in ('id','point','dish','ingredient','version'):
                    if f in r: r[f]=str(r[f]).strip()
            except DataError as e: errors.append(f'{kind}, строка {i}: {e}')
    points={str(r['id']) for r in merged['points']}
    items={str(r['id']):r for r in merged['items']}
    for kind,rows in merged.items():
        for i,r in enumerate(rows,2):
            if 'point' in r and str(r['point']) not in points: errors.append(f'{kind}:{i}: точка {r["point"]} не сопоставлена')
            for f,expected in [('dish','dish'),('ingredient','ingredient')]:
                if f not in r: continue
                item=items.get(str(r[f]))
                if not item or item.get('kind')!=expected: errors.append(f'{kind}:{i}: {f} {r[f]} не сопоставлен')
                elif f=='ingredient' and str(r.get('unit'))!=str(item['unit']): errors.append(f'{kind}:{i}: единица не совпадает с {item["unit"]}; автоматического перевода нет')
    if errors: raise DataError('\n'.join(errors[:45]))
    return tables

class Store:
    def __init__(self,path):
        Path(path).parent.mkdir(parents=True,exist_ok=True)
        self.path=Path(path)
        self.db=sqlite3.connect(path)
        self.db.execute('PRAGMA journal_mode=WAL')
        self.db.executescript('''
        CREATE TABLE IF NOT EXISTS state(key TEXT PRIMARY KEY,value TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS imports(hash TEXT PRIMARY KEY,name TEXT,created TEXT);
        CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY,created TEXT,author TEXT,kind TEXT,reason TEXT,before_json TEXT,after_json TEXT);
        CREATE TABLE IF NOT EXISTS orders(id INTEGER PRIMARY KEY,created TEXT,point TEXT,status TEXT,payload TEXT);
        ''')
        self.db.commit()
    def get(self,key,default=None):
        row=self.db.execute('SELECT value FROM state WHERE key=?',(key,)).fetchone()
        return json.loads(row[0]) if row else default
    def put(self,key,value):
        self.db.execute('INSERT OR REPLACE INTO state VALUES(?,?)',(key,json.dumps(value,ensure_ascii=False)))
    def log(self,kind,before,after,reason,author):
        self.db.execute('INSERT INTO audit(created,author,kind,reason,before_json,after_json) VALUES(?,?,?,?,?,?)',
                        (now(),author,kind,reason,json.dumps(before,ensure_ascii=False),json.dumps(after,ensure_ascii=False)))
    def import_file(self,path,kind=None,author='Продавец'):
        digest=hashlib.sha256(Path(path).read_bytes()).hexdigest()
        if self.db.execute('SELECT 1 FROM imports WHERE hash=?',(digest,)).fetchone(): return 'Этот файл уже загружен. Данные не изменены.'
        adapter=FileAdapter()
        tables=adapter.read(path,kind)
        existing=self.get('tables',{})
        was_demo=self.get('seeded_demo',False) or (bool(existing) and self.get('demo',False)!=adapter.is_demo)
        if was_demo:
            if 'points' not in tables:
                raise DataError('Первый импорт заменяет демонстрационную базу. Используйте полный XLSX или начните с points.csv, затем items.csv и зависимых таблиц.')
            existing={}
        validate(tables,existing)
        # Explicit snapshot replacement is never an additive operation.
        merged={**existing,**tables}
        with self.db:
            self.put('tables',merged)
            self.put('demo',adapter.is_demo)
            self.put('seeded_demo',False)
            self.put('import_time',now())
            self.put('revision',self.get('revision',0)+1)
            # Keep draft history but do not reuse demo edits for real data.
            if was_demo:
                self.db.execute("DELETE FROM state WHERE key LIKE 'draft:%'")
            self.db.execute('INSERT INTO imports VALUES(?,?,?)',(digest,Path(path).name,now()))
            self.log('Импорт снимка', {k:len(existing.get(k,[])) for k in tables}, {k:len(v) for k,v in tables.items()},Path(path).name,author)
        return 'Загружено: '+', '.join(f'{LABELS[k]} — {len(v)}' for k,v in tables.items())
    def seed(self):
        if self.get('tables') is not None: return
        with self.db:
            self.put('tables',demo_data())
            self.put('demo',True)
            self.put('seeded_demo',True)
            self.put('import_time',now())
            self.put('revision',1)
            self.put('settings',{'regular_share':0.6,'recent_days':28,'author':'Продавец','calculation_date':date.today().isoformat()})
    def correct_batch(self,batch_id,qty,usable,reason,author,writeoff=False):
        if not reason.strip(): raise DataError('Укажите причину')
        qty=number(qty,'Количество')
        tables=self.get('tables')
        batch=next(b for b in tables['batches'] if str(b['id'])==str(batch_id))
        before=dict(batch)
        if writeoff:
            if qty>float(batch['qty']): raise DataError('Списание превышает остаток')
            batch['qty']=float(batch['qty'])-qty
        else: batch['qty']=qty
        batch['usable']=int(usable)
        batch['checked_at']=now()
        movement={'id':f'local-{datetime.now().timestamp()}','point':batch['point'],'ingredient':batch['ingredient'],'date':date.today().isoformat(),
                  'kind':'writeoff' if writeoff else 'adjustment','qty':qty if writeoff else abs(qty-float(before['qty'])),'unit':batch['unit'],'reason':reason}
        tables.setdefault('movements',[]).append(movement)
        with self.db:
            self.put('tables',tables)
            self.put('revision',self.get('revision',0)+1)
            self.log('Списание' if writeoff else 'Сверка остатка',before,batch,reason,author)
    def save_draft(self,point,payload):
        with self.db: self.put('draft:'+point,payload)
    def confirm(self,point,payload,author):
        if payload.get('revision')!=self.get('revision'): raise DataError('Данные изменились. Пересчитайте заявку: ручные правки сохранятся.')
        if payload.get('errors'): raise DataError('Сначала устраните ошибки расчета')
        if not payload.get('stock_confirmed') or payload.get('stock_confirmed_revision')!=self.get('revision'): raise DataError('Подтвердите сверку фактических остатков')
        with self.db:
            cur=self.db.execute('INSERT INTO orders(created,point,status,payload) VALUES(?,?,?,?)',(now(),point,'Подтверждена',json.dumps(payload,ensure_ascii=False)))
            self.log('Подтверждение заявки',None,{'order':cur.lastrowid,'payload':payload},'Подтверждено продавцом',author)
        return cur.lastrowid

def demo_data(today=None):
    today=today or date.today()
    rng=random.Random(23)
    items=[]
    for id,name,price,cost,fallback in [('classic','Блин со сметаной',145,47,25),('chicken','Блин с курицей и сыром',265,104,38),('ham','Блин с ветчиной и сыром',245,91,30),('berry','Блин с вишней',205,72,18),('caviar','Блин с икрой',490,260,2)]:
        items.append(dict(id=id,name=name,kind='dish',unit='шт',price=price,cost=cost,pack=1,target=0,fallback=fallback))
    for id,name,unit,cost,pack,target in [('dough','Тесто блинное','л',95,5,3),('cream','Сметана','кг',280,1,.3),('chicken_i','Куриная начинка','кг',470,1,.5),('cheese','Сыр','кг',690,1,.4),('ham_i','Ветчина','кг',520,1,.4),('cherry','Вишневая начинка','кг',390,2,.3),('caviar_i','Икра','кг',4200,.25,.05)]:
        items.append(dict(id=id,name=name,kind='ingredient',unit=unit,price=cost,cost=cost,pack=pack,target=target,fallback=0))
    recipes=[]
    for dish,ings in {'classic':[('cream',.04)],'chicken':[('chicken_i',.08),('cheese',.025)],'ham':[('ham_i',.065),('cheese',.025)],'berry':[('cherry',.08)],'caviar':[('caviar_i',.025)]}.items():
        for ing,amount in [('dough',.12)]+ings:
            recipes.append(dict(dish=dish,ingredient=ing,amount=amount,unit='л' if ing=='dough' else 'кг',version='1',valid_from=year_start(today).isoformat()))
    sales=[]
    first=year_start(today)
    for n in range((today-first).days):
        d=first+timedelta(days=n)
        for idx,dish in enumerate([i for i in items if i['kind']=='dish']):
            available=not (n%37==0 and idx==1)
            qty=max(0,round(float(dish['fallback'])*(1.28 if d.weekday()>=5 else .96)*(1+0.12*n/365)+rng.gauss(0,5)))
            if dish['id']=='caviar': qty=rng.choice([0,0,0,0,1,2,4,6])
            sales.append(dict(date=d.isoformat(),point='central',dish=dish['id'],qty=qty if available else 0,available=int(available),closed=1))
    quantities=[14,.8,1.1,1.4,2.8,5.5,.75]
    batches=[dict(id='b'+str(j),point='central',ingredient=i['id'],qty=quantities[j],unit=i['unit'],usable=1,checked_at=now(),use_before='')
             for j,i in enumerate([i for i in items if i['kind']=='ingredient'])]
    deliveries=[dict(id='in1',point='central',ingredient='dough',qty=5,unit='л',arrival=(today+timedelta(days=1)).isoformat(),confirmed=1)]
    movements=[dict(id='m1',point='central',ingredient='cherry',date=(today-timedelta(days=5)).isoformat(),kind='writeoff',qty=.6,unit='кг',reason='Демонстрация: решение продавца')]
    return dict(points=[{'id':'central','name':'Центральная · Новосибирск'}],items=items,sales=sales,batches=batches,recipes=recipes,deliveries=deliveries,movements=movements,
                schedules=[dict(point='central',next_delivery=(today+timedelta(days=1)).isoformat(),interval_days=3)])

def history(tables,point,dish,cutoff):
    start=year_start(cutoff)
    return sorted([r for r in tables.get('sales',[]) if str(r['point'])==point and str(r['dish'])==dish and flag(r['closed'],'closed') and start<=date.fromisoformat(str(r['date']))<cutoff],key=lambda r:r['date'])

def predict(rows,targets,settings,fallback=0):
    usable=[r for r in rows if flag(r['available'],'available')]
    vals=[float(r['qty']) for r in usable]
    share=sum(v>0 for v in vals)/len(vals) if vals else 0
    segment='Регулярный' if share>=settings.get('regular_share',.6) else 'Редкий'
    window=settings.get('recent_days',28)
    recent=usable[-window:]
    avg=st.mean(vals) if vals else fallback
    def estimate(method,target,subset):
        vs=[float(r['qty']) for r in subset]
        if not vs: return fallback
        if method=='Среднее': return st.mean(vs)
        if method=='День недели':
            same=[float(r['qty']) for r in subset if date.fromisoformat(str(r['date'])).weekday()==target.weekday()][-8:]
            return st.mean(same) if same else st.mean(vs)
        if method=='Свежий спрос':
            tail=vs[-window:]; weights=[.9**(len(tail)-j-1) for j in range(len(tail))]
            return sum(v*w for v,w in zip(tail,weights))/sum(weights)
        # TSB-style probability and nonzero size smoothing, updated on observed available days only.
        probability=sum(v>0 for v in vs[:14])/min(14,len(vs))
        size=st.mean([v for v in vs[:14] if v>0]) if any(v>0 for v in vs[:14]) else fallback
        for v in vs[14:]:
            probability=.15*(v>0)+.85*probability
            if v>0: size=.15*v+.85*size
        return probability*size
    methods=['Среднее','День недели','Свежий спрос','Редкий спрос']
    scores={}
    # Strict expanding-origin one-day backtest; each origin sees only prior observations.
    for method in methods:
        errors=[]
        for j in range(max(14,len(usable)-28),len(usable)):
            target=date.fromisoformat(str(usable[j]['date']))
            prefix=[r for r in usable[:j] if date.fromisoformat(str(r['date']))>=year_start(target)]
            errors.append(abs(estimate(method,target,prefix)-float(usable[j]['qty'])))
        scores[method]=st.mean(errors) if errors else None
    allowed=['День недели','Свежий спрос'] if segment=='Регулярный' else ['Редкий спрос','Свежий спрос']
    method=min(allowed,key=lambda m:scores[m] if scores[m] is not None else math.inf) if len(vals)>=28 else 'Ручной резерв'
    expected=[max(0,estimate(method,d,usable) if method!='Ручной резерв' else fallback) for d in targets]
    spread=st.pstdev([float(r['qty']) for r in recent]) if len(recent)>1 else max(fallback*.7,1)
    total=sum(expected); delta=spread*math.sqrt(max(1,len(targets)))
    warnings=[]
    if len(vals)<28: warnings.append('Короткая история: используется резерв из настроек ассортимента')
    if len(rows)<350: warnings.append(f'Неполная история: {len(rows)} наблюдений за последние 12 месяцев')
    if len(usable)<len(rows): warnings.append(f'{len(rows)-len(usable)} дней отсутствия исключены; спрос в эти дни неизвестен')
    missing=0
    if rows:
        ds={str(r['date']) for r in rows}
        missing=(date.fromisoformat(str(rows[-1]['date']))-date.fromisoformat(str(rows[0]['date']))).days+1-len(ds)
    if missing: warnings.append(f'Пропущено {missing} дней: они не заменены нулевыми продажами')
    return dict(expected=total,low=max(0,total-delta),high=total+delta,daily=expected,segment=segment,method=method,scores=scores,observations=len(vals),warnings=warnings)

def active_recipes(tables,dish,day):
    candidates=[r for r in tables.get('recipes',[]) if str(r['dish'])==dish and date.fromisoformat(str(r['valid_from']))<=day]
    if not candidates: return []
    latest=max(str(r['valid_from']) for r in candidates)
    versions={str(r['version']) for r in candidates if str(r['valid_from'])==latest}
    if len(versions)>1: raise DataError(f'{dish}: несколько версий техкарты на одну дату')
    return [r for r in candidates if str(r['valid_from'])==latest]

def calculate(tables,point,day,settings):
    schedules=[r for r in tables.get('schedules',[]) if str(r['point'])==point]
    if not schedules: raise DataError('Нет графика поставок для этой точки')
    schedule=schedules[0]
    arrival=date.fromisoformat(str(schedule['next_delivery']))
    if arrival<day: raise DataError('Дата следующей поставки прошла. Обновите график.')
    end=arrival+timedelta(days=int(float(schedule['interval_days'])) )
    days=[day+timedelta(days=n) for n in range((end-day).days)]
    if not days or len(days)>90: raise DataError('Период покрытия должен быть от 1 до 90 дней')
    forecasts=[]; demands=defaultdict(lambda:[0.]*len(days)); contributions=defaultdict(dict); errors=[]
    for item in tables.get('items',[]):
        if item['kind']!='dish': continue
        f=predict(history(tables,point,str(item['id']),day),days,settings,float(item['fallback']))
        f.update(id=item['id'],name=item['name']); forecasts.append(f)
        for j,d in enumerate(days):
            rec=active_recipes(tables,str(item['id']),d)
            if not rec:
                errors.append(f'Нет техкарты: {item["name"]} на {d}'); continue
            for r in rec:
                ing=str(r['ingredient']); need=f['daily'][j]*float(r['amount'])
                demands[ing][j]+=need
                contributions[ing][item['name']]=contributions[ing].get(item['name'],0)+need
    result=[]
    for item in tables.get('items',[]):
        if item['kind']!='ingredient': continue
        ing=str(item['id']); demand=demands[ing]; demand_total=sum(demand)
        batches=[b for b in tables.get('batches',[]) if str(b['point'])==point and str(b['ingredient'])==ing]
        if not batches: errors.append(f'Нет фактического остатка: {item["name"]} (нужна строка даже при нуле)')
        usable=[b for b in batches if flag(b['usable'],'usable')]
        blocked=sum(float(b['qty']) for b in batches if not flag(b['usable'],'usable'))
        arrivals=[r for r in tables.get('deliveries',[]) if str(r['point'])==point and str(r['ingredient'])==ing and flag(r['confirmed'],'confirmed') and day<=date.fromisoformat(str(r['arrival']))<end]
        late=[r for r in tables.get('deliveries',[]) if str(r['point'])==point and str(r['ingredient'])==ing and flag(r['confirmed'],'confirmed') and date.fromisoformat(str(r['arrival']))<day]
        pool=[{'qty':float(b['qty']),'expires':str(b.get('use_before') or '9999-12-31')} for b in usable]
        lost=0.; shortage_before=0.; unmet=0.
        for j,d in enumerate(days):
            for b in pool:
                if b['expires']<d.isoformat(): lost+=b['qty']; b['qty']=0
            for r in arrivals:
                if str(r['arrival'])==d.isoformat(): pool.append({'qty':float(r['qty']),'expires':'9999-12-31'})
            need=demand[j]
            for b in sorted(pool,key=lambda b:b['expires']):
                used=min(b['qty'],need); b['qty']-=used; need-=used
            unmet+=need
            if d<arrival: shortage_before+=need
        stock=sum(float(b['qty']) for b in usable)
        incoming=sum(float(r['qty']) for r in arrivals)
        # Remove only unused quantities whose explicitly supplied date precedes use.
        effective=stock-lost
        raw=max(0,demand_total+float(item['target'])-effective-incoming)
        pack=float(item['pack']); packs=max(0,math.ceil(round(raw/pack,10)))
        surplus=max(0,effective+incoming-demand_total)
        waste=sum(float(m['qty']) for m in tables.get('movements',[]) if str(m['point'])==point and str(m['ingredient'])==ing and m['kind']=='writeoff' and day-timedelta(days=28)<=date.fromisoformat(str(m['date']))<day)
        # Portion equivalent is valid only if every recipe has the same ingredient usage.
        rates={float(r['amount']) for f in forecasts for r in active_recipes(tables,str(f['id']),day) if str(r['ingredient'])==ing}
        rate=next(iter(rates)) if len(rates)==1 else None
        result.append(dict(id=ing,name=item['name'],unit=item['unit'],demand=demand_total,stock=stock,effective=effective,incoming=incoming,target=float(item['target']),pack=pack,packs=packs,
            ordered=packs*pack,rounding=packs*pack-raw,surplus=surplus,blocked=blocked,expiry_excluded=lost,shortage_before=shortage_before,
            timing_gap=max(0,unmet-raw),waste=waste,cost=float(item['cost']),portion_rate=rate,contributions=contributions[ing],late=len(late),
            checked_at=min([str(b['checked_at']) for b in batches],default='Нет сверки'),unknown_dates=sum(not b.get('use_before') for b in usable)))
    closed=[str(r['date']) for r in tables.get('sales',[]) if str(r['point'])==point and flag(r['closed'],'closed') and str(r['date'])<day.isoformat()]
    return dict(point=point,date=day.isoformat(),end=end.isoformat(),arrival=arrival.isoformat(),days=len(days),rows=result,forecasts=forecasts,errors=list(dict.fromkeys(errors)),last_closed=max(closed) if closed else None,created=now())

def export_order(path,payload,order_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook(); ws=wb.active; ws.title='Заявка'
    ws.append(['Русские блины · заявка',order_id,'ДЕМО' if payload.get('demo') else 'Импортированные данные'])
    ws.append(['Точка',payload['point'],'Расчет',payload['date'],'Покрытие до',payload['end']])
    ws.append(['Смена закрыта',payload.get('last_closed'),'Автор',payload.get('author'),'Сохранено',payload.get('created')])
    ws.append(['Позиция','Ед.','Потребность','Остаток','В пути','Цель','Упаковка','Рекомендация, уп.','Итог, уп.','Итог, ед.','Стоимость','Причина правки'])
    for r in payload['rows']:
        edit=payload.get('edits',{}).get(r['id'],{})
        n=edit.get('packs',r['packs'])
        ws.append([r['name'],r['unit'],r['demand'],r['stock'],r['incoming'],r['target'],r['pack'],r['packs'],n,n*r['pack'],n*r['pack']*r['cost'],edit.get('reason','')])
    detail=wb.create_sheet('Объяснения')
    detail.append(['Позиция','Пригодный расчетный остаток','Исключено по явной дате','Не подтверждено пригодным','Излишек от округления','Дефицит до доставки','Риск избыточного запаса'])
    for r in payload['rows']: detail.append([r['name'],r['effective'],r['expiry_excluded'],r['blocked'],r['rounding'],r['shortage_before'],r['surplus']])
    models=wb.create_sheet('Прогноз')
    models.append(['Блюдо','Спрос, порций','Нижний ориентир','Верхний ориентир','Метод','Сегмент','Наблюдений','Ограничения'])
    for f in payload['forecasts']: models.append([f['name'],f['expected'],f['low'],f['high'],f['method'],f['segment'],f['observations'],'; '.join(f['warnings'])])
    meta=wb.create_sheet('Примечания')
    for line in ['Диапазон спроса — описательный ориентир, не доверительный интервал 95%.','Пригодность продуктов подтверждает продавец. Даты не выводятся из поступления.','Риск избыточного запаса не означает прогноз порчи.','Заявка сохранена локально. Отправка поставщику не выполнялась.','Все денежные суммы указаны в рублях.','Движения — отдельный журнал; остатки являются фактическим снимком и повторно на движения не корректируются.']:
        meta.append([line])
    changes=wb.create_sheet('Правки')
    changes.append(['Позиция','Упаковок','Автор','Время','Причина'])
    for id,e in payload.get('edits',{}).items(): changes.append([id,e['packs'],e.get('author'),e.get('at'),e.get('reason')])
    for sheet in wb:
        sheet.freeze_panes='A5' if sheet==ws else 'A2'
        header=4 if sheet==ws else 1
        for cell in sheet[header]: cell.fill=PatternFill('solid',fgColor='A4002A'); cell.font=Font(color='FFFFFF',bold=True)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str) and cell.value.startswith(('=','+','-','@')): cell.data_type='s'
                cell.alignment=Alignment(vertical='center',wrap_text=True)
                if isinstance(cell.value,float): cell.number_format='0.00'
        for col in sheet.columns:
            width=min(54,max(14,max(len(str(c.value or '')) for c in col)+2))
            sheet.column_dimensions[get_column_letter(col[0].column)].width=width
        sheet.auto_filter.ref=f'A{header}:{get_column_letter(sheet.max_column)}{sheet.max_row}'
    path=Path(path)
    tmp=path.with_name(path.stem+'.tmp.xlsx')
    wb.save(tmp); tmp.replace(path)

def write_templates(directory,tables=None):
    from openpyxl import Workbook
    directory=Path(directory); directory.mkdir(parents=True,exist_ok=True)
    wb=Workbook(); wb.remove(wb.active)
    for kind,fields in SCHEMAS.items():
        ws=wb.create_sheet(kind); ws.append(fields)
        for row in (tables or {}).get(kind,[]): ws.append([row.get(f,'') for f in fields])
        ws.freeze_panes='A2'
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=22
        with (directory/(kind+'.csv')).open('w',encoding='utf-8-sig',newline='') as out:
            writer=csv.DictWriter(out,fieldnames=fields+(['_dataset'] if tables else []),delimiter=';'); writer.writeheader()
            writer.writerows([{**{f:r.get(f,'') for f in fields},**({'_dataset':'DEMO'} if tables else {})} for r in (tables or {}).get(kind,[])])
    meta=wb.create_sheet('_meta'); meta.append(['dataset_type','DEMO' if tables else 'TEMPLATE']); meta.sheet_state='hidden'
    wb.save(directory/('demo.xlsx' if tables else 'import_template.xlsx'))
