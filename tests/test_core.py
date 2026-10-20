import copy
import json
from datetime import date,timedelta

import pytest
from openpyxl import load_workbook
from core import (Store,DataError,FileAdapter,demo_data,validate,calculate,predict,history,
                  export_order,write_templates,year_start,now)

@pytest.fixture
def data(): return demo_data(date.today())

@pytest.fixture
def store(tmp_path):
    s=Store(tmp_path/'test.sqlite3'); s.seed(); yield s; s.db.close()

def calc(data): return calculate(data,'central',date.today(),{'regular_share':.6,'recent_days':28})

def test_demo_schema(data): validate(data)

def test_regular_and_rare(data):
    fs={f['id']:f for f in calc(data)['forecasts']}
    assert fs['classic']['segment']=='Регулярный'
    assert fs['caviar']['segment']=='Редкий'
    assert fs['classic']['scores']['Среднее'] is not None
    assert fs['classic']['scores']['День недели'] is not None
    assert fs['classic']['low']<=fs['classic']['expected']<=fs['classic']['high']

def test_shared_ingredient_aggregated_once(data):
    result=calc(data); dough=next(r for r in result['rows'] if r['id']=='dough')
    assert dough['demand']==pytest.approx(sum(f['expected']*.12 for f in result['forecasts']))
    assert dough['stock']==14
    assert dough['portion_rate']==.12

def test_coverage_includes_lead_and_cycle(data):
    result=calc(data)
    assert result['days']==4
    assert result['end']==(date.today()+timedelta(days=4)).isoformat()

def test_rounding_and_formula(data):
    for r in calc(data)['rows']:
        raw=max(0,r['demand']+r['target']-r['effective']-r['incoming'])
        assert r['packs']*r['pack']>=raw-1e-8
        assert 0<=r['rounding']<r['pack']+1e-8
        assert r['rounding']==pytest.approx(r['packs']*r['pack']-raw)

def test_known_batch_expires_before_use(data):
    batch=data['batches'][0]; batch['use_before']=(date.today()-timedelta(days=1)).isoformat()
    r=next(r for r in calc(data)['rows'] if r['id']=='dough')
    assert r['expiry_excluded']==14
    assert r['effective']==0

def test_partial_batch_unused_after_explicit_date(data):
    batch=data['batches'][0]; batch['qty']=100; batch['use_before']=date.today().isoformat()
    r=next(r for r in calc(data)['rows'] if r['id']=='dough')
    assert 0<r['expiry_excluded']<100
    assert r['effective']<r['stock']

def test_unknown_dates_not_invented(data):
    r=next(r for r in calc(data)['rows'] if r['id']=='dough')
    assert r['expiry_excluded']==0
    assert r['unknown_dates']==1
    assert data['batches'][0]['use_before']==''

def test_seller_blocks_batch(data):
    data['batches'][0]['usable']=0
    r=next(r for r in calc(data)['rows'] if r['id']=='dough')
    assert r['stock']==0 and r['blocked']==14

def test_missing_recipe_blocks_confirmation(data,store):
    data['recipes']=[r for r in data['recipes'] if r['dish']!='classic']
    result=calc(data); assert any('Нет техкарты' in e for e in result['errors'])
    result.update(revision=1,stock_confirmed=True,stock_confirmed_revision=1)
    with pytest.raises(DataError): store.confirm('central',result,'Тест')

def test_last_closed_shift(data):
    assert calc(data)['last_closed']==(date.today()-timedelta(days=1)).isoformat()
    assert all(str(r['date'])<date.today().isoformat() for r in history(data,'central','classic',date.today()))

def test_history_strict_twelve_month_window(data):
    old=dict(data['sales'][0],date=(year_start(date.today())-timedelta(days=1)).isoformat(),qty=9999999)
    data['sales'].append(old)
    assert old not in history(data,'central','classic',date.today())

def test_future_cannot_influence_forecast(data):
    original=calc(data)['forecasts']
    data['sales'].append(dict(data['sales'][0],date=(date.today()+timedelta(days=10)).isoformat(),qty=999999))
    assert calc(data)['forecasts']==original

def test_stockouts_excluded_and_warned(data):
    f=next(f for f in calc(data)['forecasts'] if f['id']=='chicken')
    assert f['observations']<365
    assert any('отсутствия' in w for w in f['warnings'])

def test_new_dish_fallback():
    f=predict([], [date.today()+timedelta(days=n) for n in range(4)],{},7)
    assert f['expected']==28 and f['method']=='Ручной резерв'
    assert f['warnings']

def test_import_idempotent_and_real_data_not_mixed(tmp_path,store,data):
    write_templates(tmp_path,data); path=tmp_path/'demo.xlsx'
    store.save_draft('central',{'edits':{'dough':{'packs':100}}})
    store.import_file(path)
    assert store.get('demo') and not store.get('seeded_demo') and store.get('draft:central') is None
    initial=copy.deepcopy(store.get('tables')); count=store.db.execute('SELECT COUNT(*) FROM audit').fetchone()[0]
    assert 'уже' in store.import_file(path)
    assert store.get('tables')==initial
    assert store.db.execute('SELECT COUNT(*) FROM audit').fetchone()[0]==count

def test_import_atomic_failure(tmp_path,store,data):
    bad=copy.deepcopy(data); bad['batches'][0]['unit']='кг'
    write_templates(tmp_path,bad); before=store.get('tables')
    with pytest.raises(DataError): store.import_file(tmp_path/'demo.xlsx')
    assert store.get('tables')==before

def test_null_does_not_become_zero(data):
    data['batches'][0]['qty']=''
    with pytest.raises(DataError,match='не заполнено'): validate(data)

def test_duplicate_rows_rejected(data):
    data['sales'].append(dict(data['sales'][0]))
    with pytest.raises(DataError,match='дубликат'): validate(data)

@pytest.mark.parametrize('field,value',[('date','не дата'),('available',''),('closed',0)])
def test_bad_sales_fields(data,field,value):
    data['sales'][0][field]=value
    with pytest.raises(DataError): validate(data)

def test_current_shift_rejected(data):
    data['sales'][0]['date']=date.today().isoformat()
    with pytest.raises(DataError,match='сегодняшнего'): validate(data)

def test_decimal_comma_is_normalized(data):
    data['batches'][0]['qty']='12,5'; validate(data)
    assert data['batches'][0]['qty']==12.5

def test_actual_stock_correction_and_writeoff_audit(store):
    store.correct_batch('b0',10,True,'Пересчет','Анна')
    assert store.get('tables')['batches'][0]['qty']==10
    store.correct_batch('b0',2,True,'Утилизация по решению продавца','Анна',True)
    assert store.get('tables')['batches'][0]['qty']==8
    assert store.db.execute('SELECT COUNT(*) FROM audit').fetchone()[0]==2
    with pytest.raises(DataError): store.correct_batch('b0',20,True,'Ошибка','Анна',True)

def test_draft_manual_edit_survives_restart(store):
    p=calc(store.get('tables')); p['edits']={'dough':dict(packs=2,reason='Проверка',author='Анна',at=now())}
    store.save_draft('central',p)
    another=Store(store.path)
    assert another.get('draft:central')['edits']==p['edits']; another.db.close()

def test_revision_conflict_blocks_confirmation(store):
    p=calc(store.get('tables')); p.update(revision=1,stock_confirmed=True,stock_confirmed_revision=1)
    store.correct_batch('b0',12,True,'Уточнение','Анна')
    with pytest.raises(DataError,match='изменились'): store.confirm('central',p,'Анна')

def test_confirmation_and_excel(store,tmp_path):
    p=calc(store.get('tables')); p.update(revision=1,stock_confirmed=True,stock_confirmed_revision=1,demo=True,author='Анна',edits={'dough':dict(packs=0,reason='Есть резерв',author='Анна',at=now())})
    id=store.confirm('central',p,'Анна'); path=tmp_path/'order.xlsx'; export_order(path,p,id)
    wb=load_workbook(path,data_only=True)
    assert wb['Заявка']['H5'].value>0 and wb['Заявка']['I5'].value==0
    assert wb['Правки']['C2'].value=='Анна'
    assert store.db.execute('SELECT COUNT(*) FROM orders').fetchone()[0]==1
    wb.close()

def test_missing_stock_not_zero(data):
    data['batches']=data['batches'][1:]
    assert any('Нет фактического остатка' in e for e in calc(data)['errors'])

def test_unconfirmed_and_outside_window_deliveries_not_subtracted(data):
    data['deliveries'][0]['confirmed']=0
    assert next(r for r in calc(data)['rows'] if r['id']=='dough')['incoming']==0
    data['deliveries'][0]['confirmed']=1; data['deliveries'][0]['arrival']=(date.today()+timedelta(days=4)).isoformat()
    assert next(r for r in calc(data)['rows'] if r['id']=='dough')['incoming']==0

def test_versioned_recipe_changes_on_effective_day(data):
    revised=[dict(r,version='2',valid_from=(date.today()+timedelta(days=2)).isoformat(),amount=float(r['amount'])*2) for r in data['recipes'] if r['dish']=='classic']
    original=calc(data); data['recipes']+=revised; changed=calc(data)
    r1=next(r for r in original['rows'] if r['id']=='dough'); r2=next(r for r in changed['rows'] if r['id']=='dough')
    f=next(f for f in changed['forecasts'] if f['id']=='classic')
    assert r2['demand']-r1['demand']==pytest.approx(sum(f['daily'][2:])*.12)

def test_csv_utf8_semicolon(tmp_path,data):
    write_templates(tmp_path,data)
    rows=FileAdapter().read(tmp_path/'items.csv')['items']
    assert rows[0]['name']=='Блин со сметаной'

def test_real_import_clears_seed_and_is_not_marked_demo(tmp_path,store,data):
    write_templates(tmp_path,data)
    wb=load_workbook(tmp_path/'demo.xlsx'); wb['_meta']['B1']='REAL'; wb.save(tmp_path/'real.xlsx'); wb.close()
    store.import_file(tmp_path/'real.xlsx')
    assert not store.get('demo') and not store.get('seeded_demo')

def test_demo_marker_preserved_in_csv(tmp_path,data):
    write_templates(tmp_path,data)
    adapter=FileAdapter(); adapter.read(tmp_path/'items.csv')
    assert adapter.is_demo

def test_partial_demo_cannot_mix_with_real_dataset(tmp_path,store,data):
    write_templates(tmp_path,data)
    wb=load_workbook(tmp_path/'demo.xlsx'); wb['_meta']['B1']='REAL'; wb.save(tmp_path/'real.xlsx'); wb.close()
    store.import_file(tmp_path/'real.xlsx')
    before=store.get('tables')
    with pytest.raises(DataError): store.import_file(tmp_path/'items.csv','items')
    assert store.get('tables')==before and not store.get('demo')

def test_offline_engine_requires_no_network(data,monkeypatch):
    import socket
    def fail(*args,**kwargs): raise AssertionError('Network request')
    monkeypatch.setattr(socket,'socket',fail)
    assert calc(data)['rows']
